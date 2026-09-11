#!/usr/bin/env python3
"""
Rhode — PUBLICADOR DE RELATÓRIO NO GOOGLE DRIVE.

Por que existe: o conector do Drive só aceita conteúdo binário inline (base64) e o service account
NÃO tem quota de storage (o Google removeu) — então subir o .xlsx como arquivo não funciona.
O caminho que funciona é: criar uma PLANILHA GOOGLE nativa (pelo conector, na conta do dono),
compartilhar com o service account, e escrever nela via Sheets API.

O que ele transfere do .xlsx: abas, valores, FÓRMULAS (traduzidas para o locale pt-BR),
formatos numéricos, cores/fontes, larguras, painéis congelados, células mescladas,
formatação condicional e os gráficos nativos. Do .md: uma aba "Leitura" com as tabelas
markdown viradas em linhas reais.

GOTCHAS do locale pt-BR (aprendidos na marra em 31/08):
  · o separador de argumentos da fórmula é ';', não ',' — senão a célula vira #ERROR!
  · decimal em fórmula e em escala de cor usa vírgula ("0,3", não "0.3")
  · InterpolationPoint recusa valor negativo e recusa "0.0" (quer "0")
  · se um extremo da escala vira MIN/MAX, o outro extremo TEM que ser MIN/MAX também

Uso:
  python3 relatorios/_publicar_drive.py "<caminho .xlsx>" "<sheet_id>" ["<caminho .md>"]

O sheet_id vem de criar a planilha vazia no Drive (pelo conector) e compartilhá-la com
rhode-etl-936@creators-rhode.iam.gserviceaccount.com como editor.
"""
import os, re, sys, warnings
warnings.filterwarnings("ignore")
import openpyxl
from openpyxl.chart import LineChart, BarChart, PieChart
from google.oauth2 import service_account
from googleapiclient.discovery import build

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCOPES = ["https://www.googleapis.com/auth/spreadsheets",
          "https://www.googleapis.com/auth/drive"]


def obter_credenciais():
    """OAuth do dono quando disponível (única forma de CRIAR arquivo — a service
    account tem quota zero de Drive); service account como fallback para escrever
    em planilha que já exista. Ver automacao/oauth_setup.py."""
    import json as _json
    # FORCE_SA=1 → usa o service account direto (planilha criada fora do app OAuth,
    # cujo token drive.file não a enxerga; e é o caminho do cron headless).
    if os.environ.get("FORCE_SA") == "1":
        return service_account.Credentials.from_service_account_file(
            os.path.join(ROOT, "credentials.json"), scopes=SCOPES)
    raw = os.environ.get("GOOGLE_OAUTH_TOKEN")
    if not raw:
        tp = os.path.join(ROOT, "token_google.json")
        if os.path.exists(tp):
            raw = open(tp).read()
    if raw:
        from google.oauth2.credentials import Credentials
        info = _json.loads(raw)
        # respeita os escopos que o token realmente carrega (drive.file basta)
        return Credentials.from_authorized_user_info(info, info.get("scopes") or SCOPES)
    return service_account.Credentials.from_service_account_file(
        os.path.join(ROOT, "credentials.json"), scopes=SCOPES)


creds = obter_credenciais()
API = build("sheets", "v4", credentials=creds).spreadsheets()
ABA = "Leitura"
RED = {"red": 0.996, "green": 0.173, "blue": 0.333}
DARK = {"red": 0.059, "green": 0.090, "blue": 0.165}
GREY = {"red": 0.949, "green": 0.957, "blue": 0.965}
# referência do openpyxl: "'Aba'!$B$4:$B$35"  →  GridRange
REF = re.compile(r"^(?:'([^']+)'|([^!]+))!\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?$")

def formula_ptbr(v, eh_formula=True):
    """USER_ENTERED interpreta a fórmula no locale da planilha (pt_BR): o separador
    de argumentos é ';', não ','. Troca só as vírgulas FORA de string literal."""
    if not isinstance(v, str) or not v.startswith("="):
        return v
    if not eh_formula:
        # texto que só PARECE fórmula (ex.: "= CONTRIBUIÇÃO APÓS MÍDIA"): apóstrofo força texto
        return "'" + v
    out = []; dentro = False
    for i, ch in enumerate(v):
        if ch == '"':
            dentro = not dentro; out.append(ch); continue
        if dentro:
            out.append(ch); continue
        if ch == ",":
            out.append(";"); continue
        # literal numérico embutido: 0.064 -> 0,064 (decimal do locale)
        if ch == "." and i and v[i-1].isdigit() and i + 1 < len(v) and v[i+1].isdigit():
            out.append(","); continue
        out.append(ch)
    return "".join(out)

def rgb(c):
    """openpyxl color -> {red,green,blue} 0..1"""
    if c is None: return None
    v = getattr(c, "rgb", None)
    if not isinstance(v, str): return None
    v = v[-6:]
    try: r, g, b = int(v[0:2], 16), int(v[2:4], 16), int(v[4:6], 16)
    except ValueError: return None
    return {"red": r / 255, "green": g / 255, "blue": b / 255}

def conv(path, sheet_id):
    wb = openpyxl.load_workbook(path)
    meta = API.get(spreadsheetId=sheet_id).execute()
    existentes = {s["properties"]["title"]: s["properties"]["sheetId"] for s in meta["sheets"]}
    reqs = []
    # cria as abas que faltam
    for i, ws in enumerate(wb.worksheets):
        if ws.title not in existentes:
            reqs.append({"addSheet": {"properties": {
                "title": ws.title, "index": i,
                "gridProperties": {"rowCount": max(ws.max_row + 5, 30),
                                   "columnCount": max(ws.max_column + 2, 12)}}}})
    if reqs:
        API.batchUpdate(spreadsheetId=sheet_id, body={"requests": reqs}).execute()
    meta = API.get(spreadsheetId=sheet_id).execute()
    ids = {s["properties"]["title"]: s["properties"]["sheetId"] for s in meta["sheets"]}

    # 1) valores + fórmulas
    data = []
    for ws in wb.worksheets:
        linhas = []
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
            linhas.append([("" if c.value is None else formula_ptbr(c.value, c.data_type == "f"))
                           for c in row])
        data.append({"range": f"'{ws.title}'!A1", "values": linhas})
    API.values().batchUpdate(spreadsheetId=sheet_id, body={
        "valueInputOption": "USER_ENTERED", "data": data}).execute()

    # 2) formatação
    reqs = []
    for ws in wb.worksheets:
        sid = ids[ws.title]
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
            for c in row:
                fmt = {}
                if c.number_format and c.number_format != "General":
                    fmt["numberFormat"] = {"type": "NUMBER", "pattern": c.number_format}
                fill = rgb(getattr(c.fill, "fgColor", None)) if c.fill and c.fill.patternType else None
                if fill and fill != {"red": 0.0, "green": 0.0, "blue": 0.0}:
                    fmt["backgroundColor"] = fill
                tf = {}
                if c.font:
                    if c.font.bold: tf["bold"] = True
                    if c.font.size and float(c.font.size) != 11: tf["fontSize"] = int(float(c.font.size))
                    fc = rgb(c.font.color)
                    if fc: tf["foregroundColor"] = fc
                if tf: fmt["textFormat"] = tf
                if c.alignment:
                    if c.alignment.horizontal: fmt["horizontalAlignment"] = c.alignment.horizontal.upper()
                    if c.alignment.wrap_text: fmt["wrapStrategy"] = "WRAP"
                if not fmt: continue
                reqs.append({"repeatCell": {
                    "range": {"sheetId": sid, "startRowIndex": c.row - 1, "endRowIndex": c.row,
                              "startColumnIndex": c.column - 1, "endColumnIndex": c.column},
                    "cell": {"userEnteredFormat": fmt},
                    "fields": "userEnteredFormat(numberFormat,backgroundColor,textFormat,horizontalAlignment,wrapStrategy)"}})
        # larguras
        for letra, dim in ws.column_dimensions.items():
            if not dim.width: continue
            try: idx = openpyxl.utils.column_index_from_string(letra) - 1
            except Exception: continue
            reqs.append({"updateDimensionProperties": {
                "range": {"sheetId": sid, "dimension": "COLUMNS", "startIndex": idx, "endIndex": idx + 1},
                "properties": {"pixelSize": int(dim.width * 7 + 6)}, "fields": "pixelSize"}})
        # merges
        for m in ws.merged_cells.ranges:
            reqs.append({"mergeCells": {"range": {
                "sheetId": sid, "startRowIndex": m.min_row - 1, "endRowIndex": m.max_row,
                "startColumnIndex": m.min_col - 1, "endColumnIndex": m.max_col}, "mergeType": "MERGE_ALL"}})
        # freeze
        if ws.freeze_panes:
            cell = ws[ws.freeze_panes]
            reqs.append({"updateSheetProperties": {
                "properties": {"sheetId": sid, "gridProperties": {
                    "frozenRowCount": cell.row - 1, "frozenColumnCount": cell.column - 1}},
                "fields": "gridProperties(frozenRowCount,frozenColumnCount)"}})
        # formatação condicional
        for rng in ws.conditional_formatting:
            for rule in rng.rules:
                gr = []
                for sq in rng.sqref.ranges:
                    gr.append({"sheetId": sid, "startRowIndex": sq.min_row - 1, "endRowIndex": sq.max_row,
                               "startColumnIndex": sq.min_col - 1, "endColumnIndex": sq.max_col})
                if rule.type == "colorScale" and rule.colorScale:
                    cs = rule.colorScale
                    def num(v):
                        # a Sheets API recusa "0.0" (quer "0") e recusa valores negativos
                        try:
                            fv = float(v)
                        except (TypeError, ValueError):
                            return str(v), False
                        txt = str(int(fv)) if fv == int(fv) else repr(fv).replace(".", ",")
                        return txt, (fv < 0)
                    pts = []
                    tem_neg = False
                    for i, cfvo in enumerate(cs.cfvo):
                        col = rgb(cs.color[i]) or {"red": 1, "green": 1, "blue": 1}
                        t = {"num": "NUMBER", "min": "MIN", "max": "MAX",
                             "percentile": "PERCENTILE", "percent": "PERCENT"}.get(cfvo.type, "NUMBER")
                        val = None
                        if t in ("NUMBER", "PERCENT", "PERCENTILE") and cfvo.val is not None:
                            val, neg = num(cfvo.val)
                            tem_neg = tem_neg or neg
                        pts.append({"color": col, "type": t, "value": val})
                    if tem_neg:
                        # extremo negativo não é expressável: ancora a escala em MIN/MAX.
                        # (a API também exige que os dois extremos sejam do mesmo estilo)
                        pts[0] = {"color": pts[0]["color"], "type": "MIN", "value": None}
                        pts[-1] = {"color": pts[-1]["color"], "type": "MAX", "value": None}
                    def limpa(p):
                        d = {"color": p["color"], "type": p["type"]}
                        if p["value"] is not None and p["type"] in ("NUMBER", "PERCENT", "PERCENTILE"):
                            d["value"] = p["value"]
                        return d
                    grad = {"minpoint": limpa(pts[0]), "maxpoint": limpa(pts[-1])}
                    if len(pts) == 3: grad["midpoint"] = limpa(pts[1])
                    reqs.append({"addConditionalFormatRule": {
                        "rule": {"ranges": gr, "gradientRule": grad}, "index": 0}})
                elif rule.type == "dataBar" and rule.dataBar:
                    col = rgb(getattr(rule.dataBar, "color", None)) or {"red": 0.99, "green": 0.17, "blue": 0.33}
                    reqs.append({"addConditionalFormatRule": {"rule": {"ranges": gr, "gradientRule": {
                        "minpoint": {"color": {"red": 1, "green": 1, "blue": 1}, "type": "MIN"},
                        "maxpoint": {"color": col, "type": "MAX"}}}, "index": 0}})
    # remove a aba vazia padrão
    for t, s in ids.items():
        if t not in [w.title for w in wb.worksheets]:
            reqs.append({"deleteSheet": {"sheetId": s}})
    # envia em lotes
    for i in range(0, len(reqs), 400):
        API.batchUpdate(spreadsheetId=sheet_id, body={"requests": reqs[i:i + 400]}).execute()
    return len(wb.worksheets), len(reqs)



def limpa(t):
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", t)
    t = re.sub(r"__(.+?)__", r"\1", t)
    t = re.sub(r"`(.+?)`", r"\1", t)
    t = re.sub(r"\[(.+?)\]\(.+?\)", r"\1", t)
    return t.strip()

def parse(md):
    linhas = md.split("\n")
    out = []          # (tipo, celulas)
    i = 0
    while i < len(linhas):
        l = linhas[i].rstrip()
        # tabela markdown?
        if l.startswith("|") and i + 1 < len(linhas) and re.match(r"^\|[\s:\-|]+\|$", linhas[i + 1].strip()):
            cab = [limpa(c) for c in l.strip().strip("|").split("|")]
            out.append(("th", cab)); i += 2
            while i < len(linhas) and linhas[i].strip().startswith("|"):
                out.append(("td", [limpa(c) for c in linhas[i].strip().strip("|").split("|")]))
                i += 1
            out.append(("blank", [""]))
            continue
        if l.startswith("### "): out.append(("h3", [limpa(l[4:])]))
        elif l.startswith("## "): out.append(("h2", [limpa(l[3:])]))
        elif l.startswith("# "):  out.append(("h1", [limpa(l[2:])]))
        elif l.strip() in ("---", "***"): out.append(("hr", [""]))
        elif not l.strip(): out.append(("blank", [""]))
        elif l.startswith(">"): out.append(("quote", [limpa(l.lstrip("> "))]))
        else: out.append(("p", [limpa(l)]))
        i += 1
    return out

def aplicar(md_path, sheet_id):
    md = open(md_path, encoding="utf-8").read()
    blocos = parse(md)
    ncols = max(len(b[1]) for b in blocos)
    meta = API.get(spreadsheetId=sheet_id).execute()
    ids = {s["properties"]["title"]: s["properties"]["sheetId"] for s in meta["sheets"]}
    if ABA in ids:
        API.batchUpdate(spreadsheetId=sheet_id, body={"requests": [{"deleteSheet": {"sheetId": ids[ABA]}}]}).execute()
    r = API.batchUpdate(spreadsheetId=sheet_id, body={"requests": [{"addSheet": {"properties": {
        "title": ABA, "index": 0,
        "gridProperties": {"rowCount": len(blocos) + 10, "columnCount": max(ncols, 3)}}}}]}).execute()
    sid = r["replies"][0]["addSheet"]["properties"]["sheetId"]
    API.values().update(spreadsheetId=sheet_id, range=f"'{ABA}'!A1",
                        valueInputOption="RAW",
                        body={"values": [b[1] for b in blocos]}).execute()
    reqs = [
        {"updateDimensionProperties": {"range": {"sheetId": sid, "dimension": "COLUMNS", "startIndex": 0, "endIndex": 1},
                                       "properties": {"pixelSize": 560}, "fields": "pixelSize"}},
        {"updateDimensionProperties": {"range": {"sheetId": sid, "dimension": "COLUMNS", "startIndex": 1, "endIndex": max(ncols, 3)},
                                       "properties": {"pixelSize": 150}, "fields": "pixelSize"}},
        {"repeatCell": {"range": {"sheetId": sid},
                        "cell": {"userEnteredFormat": {"wrapStrategy": "WRAP",
                                                       "verticalAlignment": "TOP"}},
                        "fields": "userEnteredFormat(wrapStrategy,verticalAlignment)"}},
    ]
    for i, (tipo, cels) in enumerate(blocos):
        fmt, campos = None, None
        largura = len(cels)
        if tipo == "h1":
            fmt = {"backgroundColor": RED, "textFormat": {"bold": True, "fontSize": 15,
                   "foregroundColor": {"red": 1, "green": 1, "blue": 1}}}
        elif tipo == "h2":
            fmt = {"backgroundColor": DARK, "textFormat": {"bold": True, "fontSize": 12,
                   "foregroundColor": {"red": 1, "green": 1, "blue": 1}}}
        elif tipo == "h3":
            fmt = {"backgroundColor": GREY, "textFormat": {"bold": True, "fontSize": 11}}
        elif tipo == "th":
            fmt = {"backgroundColor": DARK, "textFormat": {"bold": True,
                   "foregroundColor": {"red": 1, "green": 1, "blue": 1}},
                   "horizontalAlignment": "CENTER"}
        elif tipo == "quote":
            fmt = {"backgroundColor": {"red": 1, "green": 0.984, "blue": 0.902},
                   "textFormat": {"italic": True}}
        if fmt:
            reqs.append({"repeatCell": {
                "range": {"sheetId": sid, "startRowIndex": i, "endRowIndex": i + 1,
                          "startColumnIndex": 0, "endColumnIndex": max(largura, 1)},
                "cell": {"userEnteredFormat": fmt},
                "fields": "userEnteredFormat(backgroundColor,textFormat,horizontalAlignment)"}})
        if tipo in ("h1", "h2", "h3", "quote", "p") and ncols > 1:
            reqs.append({"mergeCells": {"range": {"sheetId": sid, "startRowIndex": i, "endRowIndex": i + 1,
                                                  "startColumnIndex": 0, "endColumnIndex": ncols},
                                        "mergeType": "MERGE_ALL"}})
    reqs.append({"updateSheetProperties": {"properties": {"sheetId": sid,
                 "gridProperties": {"frozenRowCount": 1}}, "fields": "gridProperties.frozenRowCount"}})
    for k in range(0, len(reqs), 400):
        API.batchUpdate(spreadsheetId=sheet_id, body={"requests": reqs[k:k + 400]}).execute()
    return len(blocos)



def col_i(s):
    n = 0
    for ch in s: n = n * 26 + (ord(ch) - 64)
    return n - 1

def grid(ref, ids, aba_padrao):
    """'Aba'!$B$4:$B$35  ->  GridRange"""
    if not ref: return None
    m = REF.match(ref.strip())
    if not m: return None
    aba = m.group(1) or m.group(2) or aba_padrao
    if aba not in ids: return None
    c1, r1 = col_i(m.group(3)), int(m.group(4))
    c2 = col_i(m.group(5)) if m.group(5) else c1
    r2 = int(m.group(6)) if m.group(6) else r1
    return {"sheetId": ids[aba], "startRowIndex": r1 - 1, "endRowIndex": r2,
            "startColumnIndex": c1, "endColumnIndex": c2 + 1}

def src(gr): return {"sourceRange": {"sources": [gr]}}

def transferir(xlsx, sheet_id):
    wb = openpyxl.load_workbook(xlsx)
    meta = API.get(spreadsheetId=sheet_id).execute()
    ids = {s["properties"]["title"]: s["properties"]["sheetId"] for s in meta["sheets"]}
    # limpa gráficos já existentes (idempotente)
    limpar = [{"deleteEmbeddedObject": {"objectId": ch["chartId"]}}
              for s in meta["sheets"] for ch in s.get("charts", [])]
    if limpar:
        API.batchUpdate(spreadsheetId=sheet_id, body={"requests": limpar}).execute()
    reqs = []
    for ws in wb.worksheets:
        if ws.title not in ids: continue
        for ch in ws._charts:
            tipo = ("LINE" if isinstance(ch, LineChart) else
                    "PIE" if isinstance(ch, PieChart) else
                    ("BAR" if getattr(ch, "type", "col") == "bar" else "COLUMN") if isinstance(ch, BarChart) else None)
            if tipo is None: continue
            series = []
            for s in ch.series:
                f = getattr(getattr(s, "val", None), "numRef", None)
                gr = grid(f.f, ids, ws.title) if f is not None else None
                if gr: series.append(gr)
            cat = None
            if ch.series:
                c = ch.series[0].cat
                ref = getattr(getattr(c, "numRef", None), "f", None) or getattr(getattr(c, "strRef", None), "f", None)
                cat = grid(ref, ids, ws.title) if ref else None
            if not series: continue
            titulo = None
            try:
                t = ch.title
                if t is not None and t.tx is not None and t.tx.rich is not None:
                    titulo = "".join(r.t or "" for p in t.tx.rich.p for r in (p.r or []))
            except Exception: pass
            anc = getattr(ch, "anchor", None)
            frm = getattr(anc, "_from", None)
            pos = {"anchorCell": {"sheetId": ids[ws.title],
                                  "rowIndex": getattr(frm, "row", 0), "columnIndex": getattr(frm, "col", 0)}}
            if tipo == "PIE":
                if cat is None: continue
                spec = {"pieChart": {"legendPosition": "RIGHT_LEGEND",
                                     "domain": src(cat), "series": src(series[0])}}
            else:
                spec = {"basicChart": {"chartType": tipo, "legendPosition": "BOTTOM_LEGEND",
                                       "headerCount": 1,
                                       "domains": ([{"domain": src(cat)}] if cat else []),
                                       # barra horizontal (BAR) mede no eixo de baixo; coluna/linha, no da esquerda
                                       "series": [{"series": src(g),
                                                   "targetAxis": "BOTTOM_AXIS" if tipo == "BAR" else "LEFT_AXIS"}
                                                  for g in series]}}
            if titulo: spec["title"] = titulo
            reqs.append({"addChart": {"chart": {"spec": spec, "position": {"overlayPosition": pos}}}})
    if reqs:
        API.batchUpdate(spreadsheetId=sheet_id, body={"requests": reqs}).execute()
    return len(reqs)



if __name__ == "__main__":
    if len(sys.argv) < 3:
        print(__doc__); raise SystemExit(1)
    xlsx, sid = sys.argv[1], sys.argv[2]
    md = sys.argv[3] if len(sys.argv) > 3 else None
    n, r = conv(xlsx, sid)
    print(f"  ✓ {n} abas · {r} requests de formatação")
    g = transferir(xlsx, sid)
    print(f"  ✓ {g} gráficos nativos")
    if md and os.path.exists(md):
        k = aplicar(md, sid)
        print(f"  ✓ aba 'Leitura' com {k} linhas do .md")
    print(f"  → https://docs.google.com/spreadsheets/d/{sid}/edit")
