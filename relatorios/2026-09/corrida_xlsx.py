#!/usr/bin/env python3
"""Gerador da planilha expert do placar da Corrida de Vídeos Rhode (#CorridaRhode).
Chamado por corrida_videos_placar.py --xlsx. openpyxl: gráfico nativo, formatação condicional,
painéis congelados. Publicar no Drive com relatorios/_publicar_drive.py."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

RED = "C8102E"; DARK = "1A1A1A"; GREY = "F2F2F2"; GREEN = "1E7B34"; AMBER = "B8860B"
WHITE = "FFFFFF"; LINE = "D0D0D0"
F_HDR = Font(bold=True, color=WHITE, size=11)
F_TITLE = Font(bold=True, color=WHITE, size=16)
F_BIG = Font(bold=True, color=DARK, size=22)
F_LBL = Font(color="666666", size=9, bold=True)
F_B = Font(bold=True, color=DARK)
FILL_HDR = PatternFill("solid", fgColor=DARK)
FILL_RED = PatternFill("solid", fgColor=RED)
FILL_GREY = PatternFill("solid", fgColor=GREY)
FILL_GOLD = PatternFill("solid", fgColor="FFF3CD")
THIN = Side(style="thin", color=LINE)
BORD = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CEN = Alignment(horizontal="center", vertical="center")
LFT = Alignment(horizontal="left", vertical="center")

def _hdr(ws, row, cols, start=1):
    for j, t in enumerate(cols):
        c = ws.cell(row=row, column=start + j, value=t)
        c.font = F_HDR; c.fill = FILL_HDR; c.alignment = CEN; c.border = BORD

def build_xlsx(path, p, creators, vids, agora, ini, fim_label, tagd=None):
    wb = openpyxl.Workbook()

    # ── PAINEL ──────────────────────────────────────────────────────────
    ws = wb.active; ws.title = "Painel"; ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:F1")
    t = ws["A1"]; t.value = "🏁 CORRIDA DE VÍDEOS RHODE JEANS"; t.font = F_TITLE; t.fill = FILL_RED
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1); ws.row_dimensions[1].height = 34
    ws.merge_cells("A2:F2")
    ws["A2"] = f"#CorridaRhode · janela {ini[8:10]}/09 → {fim_label} · atualizado {agora}"
    ws["A2"].font = Font(italic=True, color="555555", size=10); ws["A2"].alignment = Alignment(indent=1)

    # cartões de KPI — a tag (TikTok, o que o grupo vê) vs o verificado (com link, base do GMV)
    tagval = f"{p['tag_total']}" if p.get("tag_total") is not None else "—"
    tagsub = (f"TikTok · informado {tagd} · faltam {p['faltam_150']} p/ 150" if p.get("tag_total") is not None
              else "conte na #CorridaRhode no TikTok e me informe")
    cards = [
        ("PUBLICAÇÕES NA TAG", tagval, tagsub),
        ("VERIFICADOS C/ LINK", f"{p['verificados']}", "base do prêmio GMV"),
        ("GMV DA CORRIDA", f"R$ {p['gmv_total']:,.0f}".replace(",", "."), f"{p['pedidos_total']} pedidos"),
        ("VIEWS", f"{p['views_total']:,}".replace(",", "."), f"{p['n_creators']} creators"),
        ("QUALIFICADAS", f"{p['n_meta1']}", f"≥5 víd cupom · {p['n_meta2']} c/ peça"),
    ]
    r0 = 4
    for i, (lbl, val, sub) in enumerate(cards):
        col = 1 + i
        ws.cell(row=r0, column=col, value=lbl).font = F_LBL
        vc = ws.cell(row=r0 + 1, column=col, value=val); vc.font = F_BIG
        ws.cell(row=r0 + 2, column=col, value=sub).font = Font(color="888888", size=9)
        ws.column_dimensions[get_column_letter(col)].width = 21

    # status das metas
    r = r0 + 4
    ws.cell(row=r, column=1, value="STATUS DAS METAS COLETIVAS").font = F_B
    metas = [
        ("Meta 3 · 150 vídeos → R$1.500 Pix Top 5 GMV", p["bateu_150"], p["faltam_150"]),
        ("Meta 4 · 500 vídeos → R$100 Pix Top 5 Volume", p["bateu_500"], p["faltam_500"]),
    ]
    for k, (lbl, ok, faltam) in enumerate(metas):
        rr = r + 1 + k
        ws.cell(row=rr, column=1, value=lbl).font = Font(size=10)
        st = ws.cell(row=rr, column=4, value="✅ BATIDA" if ok else f"faltam {faltam}")
        st.font = Font(bold=True, color=(GREEN if ok else AMBER)); st.alignment = CEN

    # dados do termômetro p/ gráfico
    rT = r + 5
    ws.cell(row=rT, column=1, value="Termômetro").font = F_LBL
    ws.cell(row=rT + 1, column=1, value="Na tag"); ws.cell(row=rT + 1, column=2, value=p["gate"])
    ws.cell(row=rT + 2, column=1, value="Meta 150"); ws.cell(row=rT + 2, column=2, value=150)
    ws.cell(row=rT + 3, column=1, value="Meta 500"); ws.cell(row=rT + 3, column=2, value=500)
    ch = BarChart(); ch.type = "bar"; ch.title = "Vídeos postados vs metas"; ch.legend = None
    ch.height = 5; ch.width = 14
    data = Reference(ws, min_col=2, min_row=rT + 1, max_row=rT + 3)
    cats = Reference(ws, min_col=1, min_row=rT + 1, max_row=rT + 3)
    ch.add_data(data, titles_from_data=False); ch.set_categories(cats)
    ws.add_chart(ch, f"D{rT}")

    if p["gmv_total"] == 0:
        ws.merge_cells(f"A{rT+6}:F{rT+6}")
        w = ws.cell(row=rT + 6, column=1,
                    value="⚠️ GMV ainda imaturo (vídeos com 1-3 dias). Ranking de dinheiro só fecha ~15/09; hoje vale o alcance.")
        w.font = Font(italic=True, color=AMBER, size=9)

    # ── TOP 5 GMV ───────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Top 5 GMV"); ws2.sheet_view.showGridLines = False
    ws2.merge_cells("A1:G1"); ws2["A1"] = "🏆 TOP 5 GMV  (desempate: mais views)"
    ws2["A1"].font = F_TITLE; ws2["A1"].fill = FILL_RED; ws2["A1"].alignment = Alignment(indent=1)
    ws2.row_dimensions[1].height = 28
    _hdr(ws2, 3, ["#", "Creator", "GMV", "Views", "Vídeos", "Pedidos", "Prêmio (se bater 150)"])
    for i, c in enumerate(p["top_gmv"], 1):
        rr = 3 + i
        vals = [i, c["creator"], c["gmv"], c["views"], c["videos"], c["pedidos"], p["premios"].get(c["creator"], "—")]
        for j, v in enumerate(vals, 1):
            cell = ws2.cell(row=rr, column=j, value=v); cell.border = BORD
            cell.alignment = LFT if j in (2, 7) else CEN
            if j == 3: cell.number_format = 'R$ #,##0'
            if j == 4: cell.number_format = '#,##0'
        if i == 1:
            for j in range(1, 8): ws2.cell(row=rr, column=j).fill = FILL_GOLD
    for col, w in zip("ABCDEFG", [4, 26, 12, 11, 8, 9, 30]): ws2.column_dimensions[col].width = w
    ws2.freeze_panes = "A4"

    # ── RANKING VOLUME ──────────────────────────────────────────────────
    ws3 = wb.create_sheet("Ranking Volume"); ws3.sheet_view.showGridLines = False
    ws3.merge_cells("A1:E1"); ws3["A1"] = "📹 TOP 5 VOLUME  (bônus R$100 se grupo bater 500)"
    ws3["A1"].font = F_TITLE; ws3["A1"].fill = FILL_RED; ws3["A1"].alignment = Alignment(indent=1)
    ws3.row_dimensions[1].height = 28
    _hdr(ws3, 3, ["#", "Creator", "Vídeos", "Views", "GMV"])
    for i, c in enumerate(p["top_vol"], 1):
        rr = 3 + i
        for j, v in enumerate([i, c["creator"], c["videos"], c["views"], c["gmv"]], 1):
            cell = ws3.cell(row=rr, column=j, value=v); cell.border = BORD
            cell.alignment = LFT if j == 2 else CEN
            if j == 4: cell.number_format = '#,##0'
            if j == 5: cell.number_format = 'R$ #,##0'
    for col, w in zip("ABCDE", [4, 26, 9, 11, 12]): ws3.column_dimensions[col].width = w
    ws3.freeze_panes = "A4"

    # ── QUALIFICAÇÃO (todas as creators) ────────────────────────────────
    ws4 = wb.create_sheet("Qualificação"); ws4.sheet_view.showGridLines = False
    ws4.merge_cells("A1:H1"); ws4["A1"] = "🎟️ QUALIFICAÇÃO POR CREATOR"
    ws4["A1"].font = F_TITLE; ws4["A1"].fill = FILL_RED; ws4["A1"].alignment = Alignment(indent=1)
    ws4.row_dimensions[1].height = 28
    _hdr(ws4, 3, ["Creator", "Vídeos", "Cupom (≥5)", "Peça (>5)", "GMV", "Views", "Pedidos", "Origem"])
    ordered = sorted(creators, key=lambda x: (-x["videos"], -x["gmv"]))
    for i, c in enumerate(ordered):
        rr = 4 + i
        vals = [c["creator"], c["videos"], "✅" if c["meta1"] else "—", "✅" if c["meta2"] else "—",
                c["gmv"], c["views"], c["pedidos"], c["origem"]]
        for j, v in enumerate(vals, 1):
            cell = ws4.cell(row=rr, column=j, value=v); cell.border = BORD
            cell.alignment = LFT if j in (1, 8) else CEN
            if j == 5: cell.number_format = 'R$ #,##0'
            if j == 6: cell.number_format = '#,##0'
    last = 3 + len(ordered)
    if last >= 4:
        ws4.conditional_formatting.add(f"B4:B{last}",
            ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color=RED))
        ws4.conditional_formatting.add(f"C4:D{last}",
            CellIsRule(operator="equal", formula=['"✅"'], fill=PatternFill("solid", fgColor="D4EDDA")))
    for col, w in zip("ABCDEFGH", [24, 8, 11, 10, 11, 10, 9, 18]): ws4.column_dimensions[col].width = w
    ws4.freeze_panes = "A4"

    # ── VÍDEOS (detalhe/auditoria) ──────────────────────────────────────
    ws5 = wb.create_sheet("Vídeos"); ws5.sheet_view.showGridLines = False
    _hdr(ws5, 1, ["Creator", "Postado em", "Views", "GMV", "Pedidos", "Legenda"])
    for i, v in enumerate(sorted(vids, key=lambda x: (x.get("post_time") or "")), 1):
        rr = 1 + i
        vals = [v["username"], (v.get("post_time") or "")[:16].replace("T", " "),
                v.get("views") or 0, round(float(v.get("gmv") or 0), 2), v.get("sku_orders") or 0,
                (v.get("title") or "")[:80]]
        for j, val in enumerate(vals, 1):
            cell = ws5.cell(row=rr, column=j, value=val)
            cell.alignment = LFT if j in (1, 2, 6) else CEN
            if j == 3: cell.number_format = '#,##0'
            if j == 4: cell.number_format = 'R$ #,##0.00'
    for col, w in zip("ABCDEF", [22, 18, 10, 12, 9, 60]): ws5.column_dimensions[col].width = w
    ws5.freeze_panes = "A2"

    wb.save(path)
